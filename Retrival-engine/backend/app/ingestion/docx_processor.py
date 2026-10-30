import re
from dataclasses import dataclass

from docx import Document


@dataclass(frozen=True)
class TextChunk:
    index: int
    content: str


def extract_docx_text(path: str) -> str:
    document = Document(path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(re.sub(r"\s+", " ", part).strip() for part in parts).strip()


def split_text(text: str, chunk_size: int, overlap: int) -> list[TextChunk]:
    if not text.strip():
        return []
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Chunk size must be positive and overlap must be smaller than chunk size")
    sections = [section.strip() for section in text.splitlines() if section.strip()]
    if len(sections) == 1:
        words = text.split()
        return [
            TextChunk(index=index, content=" ".join(words[start:start + chunk_size]))
            for index, start in enumerate(range(0, len(words), chunk_size - overlap))
            if words[start:start + chunk_size]
        ]
    target_size = min(chunk_size, 80) if len(sections) > 1 else chunk_size
    chunks: list[TextChunk] = []
    current: list[str] = []
    current_words = 0
    for section in sections:
        words = section.split()
        if len(words) > target_size:
            if current:
                chunks.append(TextChunk(index=len(chunks), content="\n".join(current).strip()))
                current = []
                current_words = 0
            stride = target_size - overlap
            for start in range(0, len(words) - target_size + 1, stride):
                chunks.append(TextChunk(
                    index=len(chunks),
                    content=" ".join(words[start:start + target_size]),
                ))
            remainder_start = ((len(words) - target_size) // stride) * stride + target_size
            current = [" ".join(words[remainder_start:])] if remainder_start < len(words) else []
            current_words = len(words) - remainder_start
            continue
        if current and current_words + len(words) > target_size:
            chunks.append(TextChunk(index=len(chunks), content="\n".join(current)))
            overlap_words = " ".join(current).split()[-overlap:] if overlap else []
            current = [" ".join(overlap_words)] if overlap_words else []
            current_words = len(overlap_words)
        current.append(section)
        current_words += len(words)
    if current:
        chunks.append(TextChunk(index=len(chunks), content="\n".join(current).strip()))
    return chunks
